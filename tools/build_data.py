"""Build the site's data bundle.

Three jobs:

1. Pull confirmed speakers out of the project's speaker tracker
   (06 Speakers/EPW2026_Speakers_Tracker.xlsx) into data/speakers.json.
   The tracker stays the single source of truth — the site never holds a
   second, hand-maintained speaker list, which is what let the 2025 portal
   drift out of sync with the agenda.

2. Bundle every data/*.json file into js/data.js as `window.EPW`.
   The bundle exists so the site works when opened straight from disk.
   Browsers block fetch() on file:// URLs, so a colleague double-clicking
   index.html would otherwise see an empty page.

3. Stamp the CSS and JS links in index.html with a content hash, so a
   deploy never serves a new page against stale cached scripts.

Usage
-----
    python tools/build_data.py            # speakers + bundle
    python tools/build_data.py --bundle   # bundle only (no tracker needed)

Run it after editing anything in data/, and after editing the tracker.

Optional tracker columns
------------------------
The script uses these if the tracker has them, and skips them if not:
  Title  — job title, shown under the name
  Bio    — paragraph for the speaker card
  Photo  — filename in images/speakers/ (defaults to a slug of the name)
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
DATA = REPO / "data"
BUNDLE = REPO / "js" / "data.js"

# The tracker lives in the project folder, not in the repo.
TRACKER = REPO.parent.parent / "06 Speakers" / "EPW2026_Speakers_Tracker.xlsx"

CONFIRMED = {"confirmed"}


def slugify(name: str) -> str:
    ascii_name = (
        unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    )
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", ascii_name).strip("-").lower()
    return re.sub(r"-{2,}", "-", slug)


def build_speakers() -> dict:
    """Read the tracker and return the speakers payload."""
    payload = {
        "_comment": (
            "GENERATED FILE — do not hand-edit. Produced by tools/build_data.py from "
            "'06 Speakers/EPW2026_Speakers_Tracker.xlsx'. Only people with Status = "
            "Confirmed are written here. Edit the tracker, then re-run: "
            "python tools/build_data.py"
        ),
        "generatedAt": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "source": "06 Speakers/EPW2026_Speakers_Tracker.xlsx",
        "people": [],
    }

    if not TRACKER.exists():
        print(f"! Tracker not found at {TRACKER}")
        print("  Writing an empty speaker list. The site will show 'speakers to be announced'.")
        return payload

    try:
        import openpyxl
    except ImportError:
        print("! openpyxl is not installed — run: pip install openpyxl")
        return payload

    wb = openpyxl.load_workbook(TRACKER, data_only=True)
    ws = wb["People"] if "People" in wb.sheetnames else wb.worksheets[0]

    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    idx = {name: i for i, name in enumerate(header)}

    def cell(row, column):
        i = idx.get(column)
        if i is None or i >= len(row) or row[i] is None:
            return ""
        return str(row[i]).strip()

    seen, skipped = {}, 0
    for row in rows:
        if not any(row):
            continue
        name = cell(row, "Name")
        if not name or name.upper() == "TBC":
            skipped += 1
            continue
        if cell(row, "Status").lower() not in CONFIRMED:
            skipped += 1
            continue

        slug = slugify(name)
        appearance = {
            "role": cell(row, "Role"),
            "component": cell(row, "Component"),
            "session": cell(row, "Session"),
        }

        if slug in seen:
            # One person can speak in several sessions — keep one card, many appearances.
            seen[slug]["appearances"].append(appearance)
            continue

        seen[slug] = {
            "id": slug,
            "name": name,
            "org": cell(row, "Organization"),
            "title": cell(row, "Title"),
            "bio": cell(row, "Bio"),
            "photo": cell(row, "Photo") or f"{slug}.webp",
            "appearances": [appearance],
        }

    payload["people"] = sorted(seen.values(), key=lambda p: p["name"].split()[-1].lower())
    print(f"  speakers: {len(payload['people'])} confirmed, {skipped} not yet confirmed or TBC")
    return payload


def bundle() -> None:
    """Concatenate data/*.json into js/data.js as window.EPW."""
    payload = {}
    for path in sorted(DATA.glob("*.json")):
        with path.open(encoding="utf-8") as fh:
            payload[path.stem] = json.load(fh)
        print(f"  bundled data/{path.name}")

    BUNDLE.parent.mkdir(parents=True, exist_ok=True)
    with BUNDLE.open("w", encoding="utf-8", newline="\n") as fh:
        fh.write("/* GENERATED FILE — do not edit.\n")
        fh.write("   Built from data/*.json by tools/build_data.py.\n")
        fh.write("   Edit the JSON in data/, then re-run: python tools/build_data.py\n")
        fh.write(f"   Built {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} */\n")
        fh.write("window.EPW = ")
        json.dump(payload, fh, ensure_ascii=False, indent=1)
        fh.write(";\n")

    size = BUNDLE.stat().st_size
    print(f"  wrote js/data.js ({size/1024:.0f} KB)")


def stamp_assets() -> None:
    """Version the CSS and JS links in index.html by content hash.

    Browsers cache these files, so after a deploy a visitor could get the new
    index.html with yesterday's main.js — the two disagree and parts of the
    page silently break. A hash in the URL changes only when the file does.
    """
    index = REPO / "index.html"
    html = index.read_text(encoding="utf-8")
    for rel in ("css/main.css", "js/data.js", "js/main.js"):
        digest = hashlib.sha1((REPO / rel).read_bytes()).hexdigest()[:8]
        html, n = re.subn(
            r'((?:href|src)=")' + re.escape(rel) + r'(?:\?v=[0-9a-f]+)?"',
            lambda m: f'{m.group(1)}{rel}?v={digest}"',
            html,
        )
        if n != 1:
            print(f"! expected one reference to {rel} in index.html, found {n}")
    with index.open("w", encoding="utf-8", newline="\n") as fh:
        fh.write(html)
    print("  stamped asset versions in index.html")


def main() -> int:
    bundle_only = "--bundle" in sys.argv

    if not bundle_only:
        speakers = build_speakers()
        with (DATA / "speakers.json").open("w", encoding="utf-8", newline="\n") as fh:
            json.dump(speakers, fh, ensure_ascii=False, indent=2)
            fh.write("\n")

    bundle()
    stamp_assets()
    print("Done.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
